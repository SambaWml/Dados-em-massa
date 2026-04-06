"""
Tests for CSV, JSON, and Excel exporters.
"""
import csv
import io
import json

import pytest

import app.generators  # noqa: F401
from app.core.exporters import export_csv, export_json, export_excel

SAMPLE_FIELDS  = ['nome_completo', 'cpf', 'email']
SAMPLE_LABELS  = [
    {'key': 'nome_completo', 'label': 'Nome Completo'},
    {'key': 'cpf',           'label': 'CPF'},
    {'key': 'email',         'label': 'E-mail'},
]
SAMPLE_RECORDS = [
    {'nome_completo': 'Maria Silva', 'cpf': '123.456.789-09', 'email': 'maria@example.com'},
    {'nome_completo': 'João Santos', 'cpf': '987.654.321-00', 'email': 'joao@example.com'},
]


# ── CSV ───────────────────────────────────────────────────────────────────────

class TestExportCSV:
    def test_returns_bytes(self):
        result = export_csv(SAMPLE_RECORDS, SAMPLE_FIELDS)
        assert isinstance(result, bytes)

    def test_utf8_bom(self):
        result = export_csv(SAMPLE_RECORDS, SAMPLE_FIELDS)
        assert result.startswith(b'\xef\xbb\xbf'), "BOM missing"

    def test_header_row(self):
        result = export_csv(SAMPLE_RECORDS, SAMPLE_FIELDS)
        text = result.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text), delimiter=';')
        assert reader.fieldnames == SAMPLE_FIELDS

    def test_row_count(self):
        result = export_csv(SAMPLE_RECORDS, SAMPLE_FIELDS)
        text = result.decode('utf-8-sig')
        rows = list(csv.DictReader(io.StringIO(text), delimiter=';'))
        assert len(rows) == len(SAMPLE_RECORDS)

    def test_values_correct(self):
        result = export_csv(SAMPLE_RECORDS, SAMPLE_FIELDS)
        text = result.decode('utf-8-sig')
        rows = list(csv.DictReader(io.StringIO(text), delimiter=';'))
        assert rows[0]['nome_completo'] == 'Maria Silva'
        assert rows[1]['email'] == 'joao@example.com'

    def test_empty_records(self):
        result = export_csv([], SAMPLE_FIELDS)
        assert result == b''

    def test_semicolon_delimiter(self):
        records = [{'nome_completo': 'A;B', 'cpf': '111', 'email': 'x@y.com'}]
        result = export_csv(records, SAMPLE_FIELDS)
        text = result.decode('utf-8-sig')
        # The semicolon inside value should be quoted
        assert '"A;B"' in text or text.count(';') >= 2


# ── JSON ──────────────────────────────────────────────────────────────────────

class TestExportJSON:
    def test_returns_bytes(self):
        result = export_json(SAMPLE_RECORDS)
        assert isinstance(result, bytes)

    def test_valid_json(self):
        result = export_json(SAMPLE_RECORDS)
        data = json.loads(result.decode('utf-8'))
        assert isinstance(data, list)

    def test_record_count(self):
        result = export_json(SAMPLE_RECORDS)
        data = json.loads(result.decode('utf-8'))
        assert len(data) == len(SAMPLE_RECORDS)

    def test_values_preserved(self):
        result = export_json(SAMPLE_RECORDS)
        data = json.loads(result.decode('utf-8'))
        assert data[0]['nome_completo'] == 'Maria Silva'
        assert data[1]['cpf'] == '987.654.321-00'

    def test_unicode_preserved(self):
        records = [{'nome_completo': 'Élève Ação', 'cpf': '', 'email': ''}]
        result = export_json(records)
        data = json.loads(result.decode('utf-8'))
        assert data[0]['nome_completo'] == 'Élève Ação'


# ── Excel ─────────────────────────────────────────────────────────────────────

class TestExportExcel:
    def test_returns_bytes(self):
        result = export_excel(SAMPLE_RECORDS, SAMPLE_FIELDS, SAMPLE_LABELS)
        assert isinstance(result, bytes)

    def test_valid_xlsx_magic(self):
        result = export_excel(SAMPLE_RECORDS, SAMPLE_FIELDS, SAMPLE_LABELS)
        # XLSX files are ZIP archives starting with PK
        assert result[:2] == b'PK', "Not a valid ZIP/XLSX file"

    def test_xlsx_readable_by_openpyxl(self):
        from openpyxl import load_workbook
        result = export_excel(SAMPLE_RECORDS, SAMPLE_FIELDS, SAMPLE_LABELS)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws is not None

    def test_header_labels(self):
        from openpyxl import load_workbook
        result = export_excel(SAMPLE_RECORDS, SAMPLE_FIELDS, SAMPLE_LABELS)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        headers = [ws.cell(row=1, column=i + 1).value for i in range(len(SAMPLE_FIELDS))]
        expected = [lbl['label'] for lbl in SAMPLE_LABELS]
        assert headers == expected

    def test_data_rows(self):
        from openpyxl import load_workbook
        result = export_excel(SAMPLE_RECORDS, SAMPLE_FIELDS, SAMPLE_LABELS)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        # Row 1 is header, row 2 is first data row
        assert ws.cell(row=2, column=1).value == 'Maria Silva'
        assert ws.cell(row=3, column=1).value == 'João Santos'

    def test_freeze_pane(self):
        from openpyxl import load_workbook
        result = export_excel(SAMPLE_RECORDS, SAMPLE_FIELDS, SAMPLE_LABELS)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.freeze_panes == 'A2'

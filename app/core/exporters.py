"""
Export generated records to CSV, JSON, or Excel.
Each function returns raw bytes ready to be sent as a file download.
"""
import csv
import io
import json
from typing import Dict, List


def export_csv(records: List[Dict], fields: List[str]) -> bytes:
    if not records:
        return b''
    output = io.StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=fields,
        delimiter=';',
        quoting=csv.QUOTE_MINIMAL,
        extrasaction='ignore',
        lineterminator='\r\n',
    )
    writer.writeheader()
    writer.writerows(records)
    # BOM so Excel opens UTF-8 correctly
    return ('\ufeff' + output.getvalue()).encode('utf-8')


def export_json(records: List[Dict]) -> bytes:
    return json.dumps(records, ensure_ascii=False, indent=2).encode('utf-8')


def export_excel(records: List[Dict], fields: List[str],
                 fields_labels: List[Dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados Gerados"

    label_map = {f['key']: f['label'] for f in fields_labels}

    # ── Header row ────────────────────────────────────────────────────────────
    hdr_font  = Font(bold=True, color='FFFFFF', size=11)
    hdr_fill  = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col_idx, fk in enumerate(fields, 1):
        cell = ws.cell(row=1, column=col_idx, value=label_map.get(fk, fk))
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align

    ws.row_dimensions[1].height = 28

    # ── Data rows ─────────────────────────────────────────────────────────────
    even_fill = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')

    for row_idx, record in enumerate(records, 2):
        for col_idx, fk in enumerate(fields, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=record.get(fk, ''))
            if row_idx % 2 == 0:
                cell.fill = even_fill

    # ── Column widths ─────────────────────────────────────────────────────────
    for col_idx, fk in enumerate(fields, 1):
        col_letter = get_column_letter(col_idx)
        header_len = len(label_map.get(fk, fk))
        sample_len = max(
            (len(str(r.get(fk, ''))) for r in records[:200]),
            default=0,
        )
        ws.column_dimensions[col_letter].width = min(max(header_len, sample_len) + 4, 45)

    # ── Freeze header ─────────────────────────────────────────────────────────
    ws.freeze_panes = 'A2'

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
